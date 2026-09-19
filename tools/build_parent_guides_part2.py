from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"
OUT.mkdir(exist_ok=True)

INK = "1F2937"
NAVY = "243B64"
BLUE = "DCE8F6"
PALE = "F3F6FA"
GOLD = "B46A22"
MUTED = "5C6777"
LINE = "D9D9D9"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, bottom=100, start=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "6")
        tag.set(qn("w:color"), LINE)


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def set_run_font(run, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)


def configure_document(doc, title, footer_label, language):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.72)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18

    title_style = styles["Title"]
    title_style.font.name = "Calibri"
    title_style.font.size = Pt(29)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.paragraph_format.space_after = Pt(12)
    title_properties = title_style._element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)

    for style_name, size, before, after in (
        ("Heading 1", 18, 15, 8),
        ("Heading 2", 13, 11, 5),
        ("Heading 3", 11, 8, 3),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.left_indent = Inches(0.28)
        style.paragraph_format.first_line_indent = Inches(-0.16)

    header = section.header.paragraphs[0]
    header.text = footer_label
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        set_run_font(run)
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.text = "Seite " if language == "de" else "Página "
    add_page_number(footer)
    for run in footer.runs:
        set_run_font(run)
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)

    props = doc.core_properties
    props.title = title
    props.subject = "NMG Weltreligionen Prüfung Teil 2"
    props.author = "Davids Weltreligionen Training"
    props.keywords = "NMG, Weltreligionen, Buddhismus, Hinduismus, Prüfung Teil 2"


def paragraph(doc, text="", bold=False, italic=False, color=None, size=None, align=None, after=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    run = p.add_run(text)
    set_run_font(run)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if size:
        run.font.size = Pt(size)
    return p


def add_rich_paragraph(doc, lead, body):
    p = doc.add_paragraph()
    first = p.add_run(lead)
    set_run_font(first)
    first.bold = True
    second = p.add_run(body)
    set_run_font(second)
    return p


def heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        for run in p.runs:
            set_run_font(run)


def numbered(doc, items):
    for index, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.22)
        p.paragraph_format.space_after = Pt(4)
        number = p.add_run(f"{index}. ")
        set_run_font(number)
        number.bold = True
        body = p.add_run(item)
        set_run_font(body)


def page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def add_table(doc, headers, rows, widths, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for idx, width in enumerate(widths):
        table.columns[idx].width = Inches(width)
        table._tbl.tblGrid.gridCol_lst[idx].set(qn("w:w"), str(int(width * 1440)))
    header = table.rows[0]
    set_repeat_header(header)
    for idx, text in enumerate(headers):
        cell = header.cells[idx]
        cell.width = Inches(widths[idx])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.text = text
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in cell.paragraphs[0].runs:
            set_run_font(run)
            run.bold = True
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor.from_string(WHITE)
    for row_idx, values in enumerate(rows):
        row = table.add_row()
        for idx, text in enumerate(values):
            cell = row.cells[idx]
            cell.width = Inches(widths[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_idx % 2:
                set_cell_shading(cell, PALE)
            cell.text = str(text)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.08
                for run in p.runs:
                    set_run_font(run)
                    run.font.size = Pt(font_size)
    paragraph(doc, "", after=2)
    return table


def add_link(doc, label, url):
    p = doc.add_paragraph()
    relationship_id = doc.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), GOLD)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend([color, underline])
    text = OxmlElement("w:t")
    text.text = label
    run.extend([properties, text])
    hyperlink.append(run)
    p._p.append(hyperlink)
    return p


DE = {
    "filename": "Praxisleitfaden_Familie_Pruefung_2_DE.docx",
    "title": "Familienleitfaden Prüfung Teil 2",
    "subtitle": "Buddhismus Hinduismus und hinduistische Gottheiten",
    "audience": "Für Davids Lernbegleitung zu Hause",
    "footer": "Weltreligionen 6. Klasse · Prüfung Teil 2 · Familienleitfaden",
    "language": "de",
    "sections": {
        "scope": "1 Was David für die Prüfung können soll",
        "buddha": "2 Buddhismus sicher erklären",
        "truths": "3 Buddhas Lehre verständlich begleiten",
        "hindu": "4 Hinduismus ohne falsche Vereinfachungen",
        "deities": "5 Hinduistische Gottheiten beschreiben",
        "plan": "6 Lernplan für zu Hause",
        "check": "7 Fragen und Kriterien für die Lernbegleitung",
    },
}


ES = {
    "filename": "Guia_familias_Prueba_2_ES.docx",
    "title": "Guía familiar Prueba 2",
    "subtitle": "Budismo hinduismo y divinidades hindúes",
    "audience": "Para acompañar el estudio de David en casa",
    "footer": "Religiones del mundo 6.º curso · Prueba 2 · Guía familiar",
    "language": "es",
    "sections": {
        "scope": "1 Qué debe saber hacer David en la prueba",
        "buddha": "2 Cómo explicar el budismo con precisión",
        "truths": "3 Cómo acompañar el estudio de las enseñanzas de Buda",
        "hindu": "4 Cómo explicar el hinduismo sin simplificaciones falsas",
        "deities": "5 Cómo describir las divinidades hindúes",
        "plan": "6 Plan de estudio en casa",
        "check": "7 Preguntas y criterios para acompañar el aprendizaje",
    },
}


def build_guide(cfg):
    de = cfg["language"] == "de"
    doc = Document()
    configure_document(doc, cfg["title"], cfg["footer"], cfg["language"])

    paragraph(doc, "", after=62)
    paragraph(doc, "NMG · 6. KLASSE · PRÜFUNG TEIL 2" if de else "NMG · 6.º CURSO · PRUEBA 2", bold=True, color=GOLD, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, after=16)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(cfg["title"])
    set_run_font(r)
    paragraph(doc, cfg["subtitle"], color=MUTED, size=15, align=WD_ALIGN_PARAGRAPH.CENTER, after=28)
    paragraph(doc, cfg["audience"], italic=True, color=INK, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, after=54)
    paragraph(doc, "Stand September 2026" if de else "Actualizado en septiembre de 2026", color=MUTED, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    page_break(doc)

    heading(doc, cfg["sections"]["scope"])
    if de:
        paragraph(doc, "Dieser Leitfaden hilft Erwachsenen, David gezielt auf den zweiten Prüfungsteil vorzubereiten. Entscheidend ist nicht das Auswendiglernen einzelner Wörter, sondern dass er Begriffe verbinden, Abläufe erzählen und Gottheiten sachlich beschreiben kann.")
        heading(doc, "Die fünf Lernbereiche", 2)
        bullets(doc, [
            "Buddhismus und Hinduismus anhand von Symbolen, Orten, Bräuchen, Schriften und Zielen vergleichen.",
            "Buddhas Lebensweg in der richtigen Reihenfolge und mit Ursache und Folge erzählen.",
            "Die vier edlen Wahrheiten sowie die Aufgabe des achtfachen Pfades erklären.",
            "Brahman, Atman, Karma, Samsara und Moksha sinnvoll miteinander verbinden.",
            "Hinduistische Gottheiten mit Name, Erkennungsmerkmal und Bedeutung beschreiben.",
        ])
        heading(doc, "So begleiten Sie wirksam", 2)
        numbered(doc, [
            "Lassen Sie David zuerst frei erklären, ohne sofort zu verbessern.",
            "Fragen Sie nach: Was bedeutet das? Woran erkennt man es? Wie hängt es zusammen?",
            "Korrigieren Sie nur einen inhaltlichen Punkt auf einmal und lassen Sie die Antwort danach neu formulieren.",
            "Beenden Sie jede Einheit mit zwei Fragen aus dem Gedächtnis und einer Frage in der Lern App.",
        ])
        add_rich_paragraph(doc, "Gute Prüfungsantworten ", "bestehen meist aus Begriff, kurzer Erklärung und einem passenden Zusammenhang oder Beispiel.")
        add_rich_paragraph(doc, "Respektvolle Sprache ", "vermeidet Aussagen wie alle Buddhisten oder alle Hindus machen dasselbe. Formulierungen wie viele, in manchen Traditionen oder je nach Region sind genauer.")
    else:
        paragraph(doc, "Esta guía ayuda a los adultos a preparar a David para la segunda parte de la prueba. Lo importante no es memorizar palabras aisladas, sino relacionar conceptos, narrar procesos y describir divinidades con precisión y respeto.")
        heading(doc, "Los cinco bloques de aprendizaje", 2)
        bullets(doc, [
            "Comparar budismo e hinduismo mediante símbolos, lugares, prácticas, textos y objetivos.",
            "Narrar la vida de Buda en el orden correcto, explicando causas y consecuencias.",
            "Explicar las cuatro nobles verdades y la función del noble camino óctuple.",
            "Relacionar Brahman, Atman, karma, samsara y moksha.",
            "Describir divinidades hindúes indicando nombre, rasgo reconocible y significado.",
        ])
        heading(doc, "Cómo acompañar eficazmente", 2)
        numbered(doc, [
            "Deje que David explique primero sin interrumpirlo.",
            "Pregunte: ¿qué significa?, ¿cómo se reconoce?, ¿cómo se relaciona con lo anterior?",
            "Corrija un solo aspecto cada vez y pídale que formule de nuevo toda la respuesta.",
            "Termine cada sesión con dos preguntas de memoria y una pregunta en la aplicación.",
        ])
        add_rich_paragraph(doc, "Una buena respuesta de examen ", "suele incluir el término, una explicación breve y una relación o ejemplo correcto.")
        add_rich_paragraph(doc, "Lenguaje respetuoso ", "evita decir que todos los budistas o todos los hindúes creen o hacen exactamente lo mismo. Es más preciso decir muchas personas, en algunas tradiciones o según la región.")
    page_break(doc)

    heading(doc, cfg["sections"]["buddha"])
    if de:
        add_table(doc, ["Station", "Was David erzählen können soll"], [
            ["Palast", "Siddhartha wächst geschützt und wohlhabend auf."],
            ["Vier Begegnungen", "Alter, Krankheit, Tod und ein Suchender zeigen ihm Leid und die Suche nach Befreiung."],
            ["Auszug", "Er verlässt sein privilegiertes Leben, um Antworten zu finden."],
            ["Askese", "Er lebt extrem streng, erkennt aber, dass Selbstquälerei nicht zum Ziel führt."],
            ["Mittlerer Weg", "Er vermeidet sowohl Luxus als auch extreme Entbehrung."],
            ["Erwachen", "Unter dem Bodhi-Baum versteht er Ursachen des Leidens und wird zum Buddha."],
            ["Lehren", "Er gibt seine Erkenntnis weiter und gründet eine Gemeinschaft."],
        ], [1.35, 5.35])
        heading(doc, "Begriffe, die sitzen müssen", 2)
        bullets(doc, [
            "Buddha bedeutet der Erwachte. Er gilt als Lehrer und nicht als Schöpfergott.",
            "Dharma bezeichnet im Buddhismus Buddhas Lehre.",
            "Sangha ist die Gemeinschaft der buddhistisch Praktizierenden, besonders die Ordensgemeinschaft.",
            "Nirvana bedeutet Befreiung von Gier, Hass und Unwissenheit und damit vom Leiden.",
            "Meditation übt Sammlung, Achtsamkeit und genaues Beobachten.",
        ])
        heading(doc, "Gute mündliche Nachfrage", 2)
        paragraph(doc, "Warum reichte die strenge Askese Siddhartha nicht? Erwartet wird: Er erkannte, dass weder Luxus noch Selbstquälerei zur befreienden Erkenntnis führen. Deshalb wählte er den Mittleren Weg.")
    else:
        add_table(doc, ["Etapa", "Lo que David debe poder narrar"], [
            ["Palacio", "Siddhartha crece protegido y con privilegios."],
            ["Cuatro encuentros", "Ve vejez, enfermedad, muerte y a un buscador; descubre el sufrimiento y la búsqueda de liberación."],
            ["Salida", "Abandona su vida privilegiada para buscar respuestas."],
            ["Ascetismo", "Vive con extrema austeridad, pero comprende que mortificarse no conduce a la meta."],
            ["Camino medio", "Evita tanto el lujo como la privación extrema."],
            ["Despertar", "Bajo el árbol Bodhi comprende las causas del sufrimiento y se convierte en Buda."],
            ["Enseñanza", "Comparte lo comprendido y forma una comunidad."],
        ], [1.35, 5.35])
        heading(doc, "Conceptos imprescindibles", 2)
        bullets(doc, [
            "Buda significa el Despierto. Se considera un maestro, no un dios creador.",
            "Dharma se refiere en el budismo a la enseñanza de Buda.",
            "Sangha es la comunidad budista, especialmente la comunidad monástica.",
            "Nirvana es la liberación de la codicia, el odio y la ignorancia y, con ello, del sufrimiento.",
            "La meditación ejercita concentración, atención consciente y observación.",
        ])
        heading(doc, "Una buena pregunta oral", 2)
        paragraph(doc, "¿Por qué no le bastó a Siddhartha el ascetismo extremo? Respuesta esperada: comprendió que ni el lujo ni la mortificación conducen al despertar; por eso eligió el camino medio.")
    page_break(doc)

    heading(doc, cfg["sections"]["truths"])
    if de:
        add_table(doc, ["Schritt", "Kernaussage", "Leitfrage"], [
            ["1 Beobachtung", "Leid und Unzufriedenheit gehören zum Leben.", "Was ist das Problem?"],
            ["2 Ursache", "Gier, Hass und Unwissenheit verursachen Leiden.", "Warum entsteht es?"],
            ["3 Ziel", "Wenn die Ursachen enden, kann Leiden enden.", "Was ist möglich?"],
            ["4 Weg", "Der achtfache Pfad zeigt den praktischen Weg.", "Wie gelingt es?"],
        ], [1.25, 3.75, 1.7])
        heading(doc, "Der achtfache Pfad", 2)
        paragraph(doc, "Rechte Sicht, rechte Absicht, rechte Rede, rechtes Handeln, rechter Lebensunterhalt, rechtes Bemühen, rechte Achtsamkeit und rechte Sammlung. Für die Prüfung sollte David mindestens vier Teile nennen und erklären können, dass der Pfad Einsicht, ethisches Handeln und geistige Übung verbindet.")
        heading(doc, "Häufige Fehler", 2)
        add_table(doc, ["Unvollständig oder falsch", "Besser"], [
            ["Buddhisten glauben an Buddha als Schöpfergott.", "Buddha ist der erwachte Lehrer; ein Schöpfergott steht nicht im Zentrum."],
            ["Nirvana ist ein Himmel oder ein Ort.", "Nirvana bezeichnet Befreiung vom Leiden und seinen Ursachen."],
            ["Die erste Wahrheit sagt, dass nur Leid existiert.", "Sie erkennt an, dass das Leben auch Leid und Unzufriedenheit enthält."],
            ["Der Pfad ist eine Abfolge von acht Stufen.", "Die acht Bereiche werden gemeinsam geübt und verstärken einander."],
        ], [2.7, 4.0])
        add_rich_paragraph(doc, "Antwortmodell ", "Die zweite edle Wahrheit erklärt die Ursache des Leidens. Gier, Hass und Unwissenheit binden Menschen an Unzufriedenheit. Der achtfache Pfad übt eine andere Sicht und ein hilfreiches Handeln.")
    else:
        add_table(doc, ["Paso", "Idea central", "Pregunta guía"], [
            ["1 Observación", "El sufrimiento y la insatisfacción forman parte de la vida.", "¿Cuál es el problema?"],
            ["2 Causa", "Codicia, odio e ignorancia producen sufrimiento.", "¿Por qué aparece?"],
            ["3 Meta", "Si terminan las causas, puede terminar el sufrimiento.", "¿Qué es posible?"],
            ["4 Camino", "El noble camino óctuple muestra la práctica.", "¿Cómo se consigue?"],
        ], [1.25, 3.75, 1.7])
        heading(doc, "El noble camino óctuple", 2)
        paragraph(doc, "Comprensión correcta, intención correcta, palabra correcta, acción correcta, modo de vida correcto, esfuerzo correcto, atención correcta y concentración correcta. David debería nombrar al menos cuatro y explicar que el camino une comprensión, conducta ética y entrenamiento mental.")
        heading(doc, "Errores frecuentes", 2)
        add_table(doc, ["Incompleto o incorrecto", "Mejor formulación"], [
            ["Los budistas creen en Buda como dios creador.", "Buda es el maestro despierto; un dios creador no ocupa el centro."],
            ["El nirvana es un cielo o un lugar.", "Nirvana designa la liberación del sufrimiento y de sus causas."],
            ["La primera verdad dice que solo existe sufrimiento.", "Reconoce que la vida también contiene sufrimiento e insatisfacción."],
            ["El camino es una escalera de ocho pasos sucesivos.", "Sus ocho ámbitos se practican conjuntamente y se refuerzan."],
        ], [2.7, 4.0])
        add_rich_paragraph(doc, "Modelo de respuesta ", "La segunda noble verdad explica la causa del sufrimiento. La codicia, el odio y la ignorancia mantienen la insatisfacción. El noble camino óctuple entrena otra forma de comprender y actuar.")
    page_break(doc)

    heading(doc, cfg["sections"]["hindu"])
    if de:
        add_table(doc, ["Begriff", "Prüfungssichere Erklärung"], [
            ["Brahman", "Die höchste, alles durchdringende göttliche Wirklichkeit. Nicht mit Brahma verwechseln."],
            ["Atman", "Das innerste Selbst oder die Seele eines Lebewesens."],
            ["Karma", "Handlungen haben Folgen für dieses und weitere Leben."],
            ["Samsara", "Kreislauf von Geburt, Tod und Wiedergeburt."],
            ["Moksha", "Befreiung aus Samsara."],
            ["Dharma", "Pflicht, verantwortliches Handeln und Ordnung des Lebens."],
            ["Puja", "Verehrungszeremonie mit Gebeten und Gaben, im Mandir oder am Hausaltar."],
            ["Murti", "Geweihte Darstellung einer Gottheit, die bei der Verehrung gegenwärtig gedacht wird."],
        ], [1.35, 5.35])
        heading(doc, "Der wichtigste Zusammenhang", 2)
        paragraph(doc, "Eine Handlung prägt das Karma. Karma beeinflusst weitere Wiedergeburten im Samsara. Moksha ist die Befreiung aus diesem Kreislauf. Diese Begriffe sollten nicht nur einzeln definiert, sondern als Kette erklärt werden.")
        heading(doc, "Vielfalt korrekt ausdrücken", 2)
        bullets(doc, [
            "Der Hinduismus hat keinen einzelnen Gründer und entwickelte sich über lange Zeit in Südasien.",
            "Viele Gottheiten können als unterschiedliche Formen oder Zugänge zur göttlichen Wirklichkeit verstanden werden.",
            "Familien und Traditionen können verschiedene Gottheiten, Feste und Rituale in den Mittelpunkt stellen.",
            "Viele Hindus leben vegetarisch, aber nicht alle. Regionale Praxis und persönliche Entscheidungen unterscheiden sich.",
        ])
        add_rich_paragraph(doc, "Wichtige Verwechslung ", "Brahman ist die umfassende göttliche Wirklichkeit. Brahma ist eine Gottheit, die mit Schöpfung verbunden wird.")
    else:
        add_table(doc, ["Concepto", "Explicación adecuada para la prueba"], [
            ["Brahman", "La realidad divina suprema que lo impregna todo. No confundir con Brahma."],
            ["Atman", "El yo más profundo o alma de un ser vivo."],
            ["Karma", "Las acciones tienen consecuencias en esta vida y en otras vidas."],
            ["Samsara", "Ciclo de nacimiento, muerte y renacimiento."],
            ["Moksha", "Liberación del samsara."],
            ["Dharma", "Deber, conducta responsable y orden de la vida."],
            ["Puja", "Ceremonia de veneración con oraciones y ofrendas, en el mandir o en el altar doméstico."],
            ["Murti", "Representación consagrada de una divinidad, entendida como presente durante la veneración."],
        ], [1.35, 5.35])
        heading(doc, "La relación más importante", 2)
        paragraph(doc, "Una acción influye en el karma. El karma influye en futuros renacimientos dentro del samsara. Moksha es la liberación de ese ciclo. David debe explicar estos conceptos como una cadena y no solo definirlos por separado.")
        heading(doc, "Cómo expresar correctamente la diversidad", 2)
        bullets(doc, [
            "El hinduismo no tiene un único fundador y se desarrolló durante mucho tiempo en el sur de Asia.",
            "Muchas divinidades pueden entenderse como formas o caminos de acceso a la realidad divina.",
            "Familias y tradiciones pueden dar importancia a distintas divinidades, fiestas y rituales.",
            "Muchos hindúes son vegetarianos, pero no todos; la práctica varía según región y persona.",
        ])
        add_rich_paragraph(doc, "Confusión importante ", "Brahman es la realidad divina que todo lo abarca. Brahma es una divinidad asociada con la creación.")
    page_break(doc)

    heading(doc, cfg["sections"]["deities"])
    if de:
        add_table(doc, ["Name", "Erkennungsmerkmal", "Bedeutung oder Aufgabe"], [
            ["Brahma", "oft vier Köpfe", "Schöpfung"],
            ["Vishnu", "Krone und Diskus; Avatare Rama und Krishna", "Bewahrung und Schutz"],
            ["Shiva", "Dreizack oder Tanz", "Verwandlung, Zerstörung und Neubeginn"],
            ["Ganesha", "Elefantenkopf", "Neuanfänge und Hindernisse überwinden"],
            ["Lakshmi", "Lotus", "Glück, Wohlstand und Segen"],
            ["Saraswati", "Buch oder Musikinstrument", "Wissen, Sprache, Musik und Kunst"],
        ], [1.15, 2.65, 2.9])
        heading(doc, "Dreiteiliger Antwortbauplan", 2)
        numbered(doc, ["Name nennen.", "Ein sichtbares Erkennungsmerkmal beschreiben.", "Bedeutung oder Aufgabe erklären."])
        add_rich_paragraph(doc, "Beispiel Ganesha ", "Ganesha erkennt man an seinem Elefantenkopf. Er wird besonders mit Neuanfängen verbunden und gilt als Helfer beim Überwinden von Hindernissen.")
        add_rich_paragraph(doc, "Beispiel Vishnu ", "Vishnu gilt als Bewahrer und Beschützer. Rama und Krishna werden in vielen Traditionen als seine Avatare verstanden.")
        heading(doc, "So üben Sie mit Bildern", 2)
        numbered(doc, [
            "Zeigen Sie das Bild zunächst ohne Namen und lassen Sie nur sichtbare Merkmale beschreiben.",
            "David nennt die Gottheit und begründet die Zuordnung mit einem Merkmal.",
            "Er ergänzt die Bedeutung in einem vollständigen Satz.",
            "Bei Unsicherheit wird nur das Merkmal wiederholt; danach versucht David die ganze Antwort erneut.",
        ])
        add_rich_paragraph(doc, "Nicht bewerten ", "Götterbilder sind religiöse Darstellungen. Beschrieben werden Name, Symbole und Bedeutung, nicht ob eine Religion richtig oder falsch ist.")
    else:
        add_table(doc, ["Nombre", "Rasgo reconocible", "Significado o función"], [
            ["Brahma", "a menudo cuatro cabezas", "creación"],
            ["Vishnu", "corona y disco; avatares Rama y Krishna", "conservación y protección"],
            ["Shiva", "tridente o danza", "transformación, destrucción y nuevo comienzo"],
            ["Ganesha", "cabeza de elefante", "nuevos comienzos y superación de obstáculos"],
            ["Lakshmi", "loto", "fortuna, prosperidad y bendición"],
            ["Saraswati", "libro o instrumento musical", "conocimiento, lenguaje, música y arte"],
        ], [1.15, 2.65, 2.9])
        heading(doc, "Estructura de respuesta en tres partes", 2)
        numbered(doc, ["Nombrar la divinidad.", "Describir un rasgo visible.", "Explicar su significado o función."])
        add_rich_paragraph(doc, "Ejemplo Ganesha ", "Ganesha se reconoce por su cabeza de elefante. Se asocia especialmente con los nuevos comienzos y con la superación de obstáculos.")
        add_rich_paragraph(doc, "Ejemplo Vishnu ", "Vishnu se considera conservador y protector. En muchas tradiciones Rama y Krishna se entienden como avatares suyos.")
        heading(doc, "Cómo practicar con imágenes", 2)
        numbered(doc, [
            "Muestre la imagen sin el nombre y pida primero una descripción de lo visible.",
            "David nombra la divinidad y justifica la identificación mediante un rasgo.",
            "Añade el significado en una frase completa.",
            "Si duda, se repite solo el rasgo; después David formula de nuevo la respuesta completa.",
        ])
        add_rich_paragraph(doc, "No emitir juicios ", "Las imágenes son representaciones religiosas. Se describen nombre, símbolos y significado, sin valorar si una religión es verdadera o falsa.")
    page_break(doc)

    heading(doc, cfg["sections"]["plan"])
    if de:
        add_table(doc, ["Tag", "Schwerpunkt", "15 bis 25 Minuten"], [
            ["1", "Überblick", "Vergleichskarte lesen; je drei Merkmale beider Religionen frei nennen; 10 App Fragen."],
            ["2", "Buddhas Leben", "Sieben Stationen ordnen und laut erzählen; Ursache und Folge nachfragen; 10 App Fragen."],
            ["3", "Buddhas Lehre", "Vier Wahrheiten als Problem Ursache Ziel Weg erklären; vier Teile des Pfades nennen."],
            ["4", "Hinduismus", "Brahman bis Moksha erklären; Karma Samsara Moksha als Kette formulieren; 10 App Fragen."],
            ["5", "Gottheiten", "Sechs Bilder zuordnen; jede Gottheit nach Name Merkmal Bedeutung beschreiben."],
            ["6", "Vergleichen", "Nirvana und Moksha, Dharma in beiden Religionen sowie Tempel und Bräuche vergleichen."],
            ["7", "Probe", "Zehn Fragen ohne Hilfe; Fehler notieren; in der App gezielt die schwachen Themen wiederholen."],
        ], [0.48, 1.45, 4.77], font_size=8.7)
        heading(doc, "Eine kurze Einheit", 2)
        numbered(doc, [
            "3 Minuten: Zwei Begriffe vom Vortag aus dem Gedächtnis erklären.",
            "7 Minuten: Einen neuen Abschnitt lesen und mit eigenen Worten zusammenfassen.",
            "8 Minuten: Fünf bis zehn Fragen in der Lern App beantworten.",
            "4 Minuten: Zwei Fehler richtig erklären und einen Merksatz formulieren.",
        ])
        add_rich_paragraph(doc, "Wann abbrechen ", "Wenn Antworten zunehmend geraten werden, ist eine Pause wirksamer als weiteres Wiederholen. Lieber am nächsten Tag kurz und gezielt fortsetzen.")
        add_rich_paragraph(doc, "App sinnvoll nutzen ", "Im Elternbereich ist Prüfung Teil 2 standardmässig ausgewählt. Achten Sie besonders auf Themen mit niedriger Genauigkeit; schwierige Fragen werden in neuen Runden häufiger angeboten.")
    else:
        add_table(doc, ["Día", "Enfoque", "15 a 25 minutos"], [
            ["1", "Visión general", "Leer la comparación; decir tres rasgos de cada religión; 10 preguntas en la aplicación."],
            ["2", "Vida de Buda", "Ordenar siete etapas y narrarlas; preguntar por causa y consecuencia; 10 preguntas."],
            ["3", "Enseñanza de Buda", "Explicar las verdades como problema causa meta camino; nombrar cuatro aspectos del camino."],
            ["4", "Hinduismo", "Explicar de Brahman a moksha; formular la cadena karma samsara moksha; 10 preguntas."],
            ["5", "Divinidades", "Identificar seis imágenes; describir cada una con nombre rasgo significado."],
            ["6", "Comparación", "Comparar nirvana y moksha, dharma en ambas religiones, templos y prácticas."],
            ["7", "Simulación", "Diez preguntas sin ayuda; anotar errores; repetir en la aplicación los temas débiles."],
        ], [0.48, 1.45, 4.77], font_size=8.7)
        heading(doc, "Una sesión breve", 2)
        numbered(doc, [
            "3 minutos: explicar de memoria dos conceptos del día anterior.",
            "7 minutos: leer una sección nueva y resumirla con palabras propias.",
            "8 minutos: responder de cinco a diez preguntas en la aplicación.",
            "4 minutos: explicar correctamente dos errores y formular una frase clave.",
        ])
        add_rich_paragraph(doc, "Cuándo parar ", "Si las respuestas se convierten en intentos al azar, una pausa es más útil que seguir repitiendo. Es mejor continuar al día siguiente de forma breve y dirigida.")
        add_rich_paragraph(doc, "Uso útil de la aplicación ", "En el panel parental aparece seleccionada por defecto la Prueba 2. Conviene observar los temas con menor precisión; las preguntas difíciles reaparecen con mayor frecuencia en nuevas rondas.")
    page_break(doc)

    heading(doc, cfg["sections"]["check"])
    if de:
        heading(doc, "Zehn mündliche Kontrollfragen", 2)
        numbered(doc, [
            "Warum verliess Siddhartha sein geschütztes Leben?",
            "Was bedeutet der Mittlere Weg?",
            "Wie führen die vier edlen Wahrheiten vom Problem zum Lösungsweg?",
            "Nenne vier Bereiche des achtfachen Pfades und erkläre seine Aufgabe.",
            "Wie hängen Karma, Samsara und Moksha zusammen?",
            "Was ist der Unterschied zwischen Brahman und Brahma?",
            "Was geschieht bei einer Puja und wo kann sie stattfinden?",
            "Beschreibe Ganesha nach Name, Merkmal und Bedeutung.",
            "Nenne je einen religiösen Ort und einen Brauch beider Religionen.",
            "Vergleiche Nirvana und Moksha in zwei vollständigen Sätzen.",
        ])
        heading(doc, "Woran Sie eine gute Antwort erkennen", 2)
        add_table(doc, ["Kriterium", "Erfüllt, wenn"], [
            ["Richtig", "Begriffe und Zusammenhänge stimmen sachlich."],
            ["Vollständig", "Die Antwort enthält Erklärung und nicht nur ein Stichwort."],
            ["Verbunden", "Ursache, Folge oder Vergleich wird ausdrücklich genannt."],
            ["Eigene Worte", "David kann erklären, ohne einen Satz auswendig abzulesen."],
            ["Respektvoll", "Vielfalt wird anerkannt; pauschale Urteile werden vermieden."],
        ], [1.35, 5.35])
        heading(doc, "Zusatzmaterial", 2)
        paragraph(doc, "Die beiden SRF Kids Videos eignen sich zur Wiederholung. Beim zweiten Sehen sollte David drei Schlüsselbegriffe notieren und danach eine Minute frei zusammenfassen.")
        add_link(doc, "SRF Kids Die Weltreligion Buddhismus", "https://www.srf.ch/play/tv/srf-kids---clip-und-klar/video/die-weltreligion-buddhismus?urn=urn%3Asrf%3Avideo%3A108463a7-a364-4397-906a-51f07867d885")
        add_link(doc, "SRF Kids Die Weltreligion Hinduismus", "https://www.srf.ch/play/tv/srf-kids---clip-und-klar/video/die-weltreligion-hinduismus?urn=urn%3Asrf%3Avideo%3Acb5925cd-6358-4659-bca5-e96ad0b0ddd1")
    else:
        heading(doc, "Diez preguntas orales de control", 2)
        numbered(doc, [
            "¿Por qué abandonó Siddhartha su vida protegida?",
            "¿Qué significa el camino medio?",
            "¿Cómo llevan las cuatro nobles verdades del problema al camino de solución?",
            "Nombra cuatro aspectos del noble camino óctuple y explica para qué sirve.",
            "¿Cómo se relacionan karma, samsara y moksha?",
            "¿Cuál es la diferencia entre Brahman y Brahma?",
            "¿Qué ocurre en una puja y dónde puede celebrarse?",
            "Describe a Ganesha mediante nombre, rasgo y significado.",
            "Nombra un lugar religioso y una práctica de cada religión.",
            "Compara nirvana y moksha en dos frases completas.",
        ])
        heading(doc, "Cómo reconocer una buena respuesta", 2)
        add_table(doc, ["Criterio", "Se cumple cuando"], [
            ["Correcta", "Los conceptos y relaciones son exactos."],
            ["Completa", "Incluye una explicación y no solo una palabra clave."],
            ["Relacionada", "Expresa una causa, una consecuencia o una comparación."],
            ["Palabras propias", "David puede explicarlo sin leer una frase memorizada."],
            ["Respetuosa", "Reconoce la diversidad y evita juicios generales."],
        ], [1.35, 5.35])
        heading(doc, "Material adicional", 2)
        paragraph(doc, "Los dos vídeos de SRF Kids sirven para repasar. En el segundo visionado, David debería anotar tres conceptos clave y después hacer un resumen oral de un minuto.")
        add_link(doc, "SRF Kids La religión mundial del budismo", "https://www.srf.ch/play/tv/srf-kids---clip-und-klar/video/die-weltreligion-buddhismus?urn=urn%3Asrf%3Avideo%3A108463a7-a364-4397-906a-51f07867d885")
        add_link(doc, "SRF Kids La religión mundial del hinduismo", "https://www.srf.ch/play/tv/srf-kids---clip-und-klar/video/die-weltreligion-hinduismus?urn=urn%3Asrf%3Avideo%3Acb5925cd-6358-4659-bca5-e96ad0b0ddd1")

    output = OUT / cfg["filename"]
    doc.save(output)
    return output


if __name__ == "__main__":
    for config in (DE, ES):
        print(build_guide(config))
